import json,os,sys,urllib.request
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
T=os.environ['TEMP']+'/pm/'
nm=json.load(open(T+'newmodels.json')); f=sys.argv[1]; wb=load_workbook(f)
ws=wb['Master catalogue']
# add search-count columns
hdr=[c.value for c in ws[1]]
if 'Papers in new search' not in hdr:
    col=ws.max_column+1
    for i,h in enumerate(['Papers in new search','Title mentions','Papers 2024-2026']):
        ws.cell(1,col+i,h).font=ws.cell(1,1).font.copy(); ws.cell(1,col+i).fill=ws.cell(1,1).fill.copy()
else: col=hdr.index('Papers in new search')+1
bym={r[0]:r for r in nm['catalogue']}
for row in ws.iter_rows(min_row=2):
    m=row[1].value
    key=next((k for k in bym if m.startswith(k[:20])),None)
    if key:
        r=bym[key]
        for i,v in enumerate(r[2:5]): ws.cell(row[0].row,col+i,v)
    if m and m.startswith('CAM / RCAM'): row[12].value='Search count inflated by CAM-ICU (delirium); true count much lower. '+(row[12].value or '')
    if m and ('prediction rules' in m or 'progression models' in m or 'SFTS' in m or 'Leptospirosis' in m or 'biomarker triage' in m):
        row[12].value='Count = prediction papers on this disease, not one named model. '+(row[12].value or '')
# candidate originator OA status
pm=[r[9].split()[0] for r in nm['new'] if r[9]]
d=json.load(urllib.request.urlopen('https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?db=pubmed&retmode=json&id='+','.join(pm)))['result']
pmc={p:next((a['value'] for a in d.get(p,{}).get('articleids',[]) if a['idtype']=='pmc'),'') for p in pm}
NEWF=PatternFill('solid',fgColor='DDEBF7')
for r in nm['new']:
    dom,name,typ,popn,status,rx,a,t,rec,early=r
    p=early.split()[0] if early else ''
    ft=(f"{p}: {'PMC free' if pmc.get(p) else 'NEED PDF'}") if p else ''
    vals=[dom,name,typ,popn,status,0,'NEW (search)',p,'',ft,'','','Added from 2026 search. Originator = earliest paper with name in title; VERIFY.']
    ws.append(vals+[None]*(col-1-len(vals))+[a,t,rec])
    for c in ws[ws.max_row]: c.fill=NEWF
wb.save(f); print('added',len(nm['new']),'new models')
