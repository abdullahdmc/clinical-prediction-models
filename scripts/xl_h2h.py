import json,os,csv,sys,time,urllib.request
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
T=os.environ['TEMP']+'/pm/'
ids=json.load(open(T+'serial_ids.json'))
rows=list(csv.DictReader(open(T+'h2h_extract.tsv',encoding='utf-8'),delimiter='\t'))
# Abstracts state scores were compared by AUC but give no values -> full text needed
FT={'7':'SOFA/qSOFA comparators for ML model','137':'MEDS and APACHE II AUCs','146':'APACHE II, SOFA, qSOFA AUCs','296':'SOFA AUC for SA-AKI',
'564':'single-score AUCs for qSOFA, MEWS, NEWS','643':'SAPS III, MPM0-III, APACHE II, SOFA AUCs','754':'MEWS and SAPS II AUCs','806':'APACHE II AUC',
'941':'MEDS AUC','1033':'SOFA and SAPS II AUCs','1050':'APACHE II and SOFA AUCs','1074':'SOFA and APACHE II AUCs (stated, values omitted)',
'1079':'qSOFA, SIRS, SOFA, APACHE II, SAPS II/III, NEWS, MEDS AUCs (rich comparison)','1081':'APACHE II, SAPS II, SOFA AUCs','1101':'SOFA AUC',
'1311':'SOFA, APACHE II AUCs','1313':'APACHE II, SOFA AUCs','1332':'APACHE II, MEDS, SAPS II AUCs by sepsis source','1401':'SOFA and APACHE II AUCs',
'1467':'APACHE II, MEDS, SOFA AUCs','1480':'MEDS and APACHE II AUCs (n=859)','925':'APACHE II and SOFA AUCs','830':'SOFA AUC','196':'SOFA AUC','777':'SOFA AUC'}
pm=[ids[s][0] for s in FT]
pmc={}
d=json.load(urllib.request.urlopen('https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?db=pubmed&retmode=json&id='+','.join(pm)))['result']
for p in pm: pmc[p]=next((a['value'] for a in d.get(p,{}).get('articleids',[]) if a['idtype']=='pmc'),'')
f=sys.argv[1]; wb=load_workbook(f)
for n in ('Head-to-head results','Comparison full text'):
    if n in wb.sheetnames: del wb[n]
H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78'); DER=PatternFill('solid',fgColor='FFF2CC'); COMB=PatternFill('solid',fgColor='E2EFDA')
def sheet(ws,head,data,widths):
    ws.append(head)
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='top')
    for r in data: ws.append(r)
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
ws=wb.create_sheet('Head-to-head results')
data=[]
for r in rows:
    p=ids[r['serial']][0]
    kind='Score + biomarker / new model' if ('+' in r['model'] or 'nomogram' in r['model'].lower() or 'ML' in r['model'] or 'SSMART' in r['model'] or 'biomarker' in r['model'].lower()) else 'Established score'
    data.append([int(r['serial']),r['author'],int(r['year']),p,r['setting'],r['age'],int(r['n']),r['outcome'],r['model'],kind,float(r['auc']),r['ci'],r['value_type'],r['note'],f'https://pubmed.ncbi.nlm.nih.gov/{p}/'])
data.sort(key=lambda x:(x[0],x[9]!='Established score'))
sheet(ws,['Serial (your file)','First author','Year','PMID','Setting','Population','N','Outcome','Model','Model class','AUC','95% CI','Value type','Notes','PubMed'],data,[8,14,6,10,18,16,7,26,34,18,7,12,10,40,30])
for row in ws.iter_rows(min_row=2):
    fill=DER if row[12].value=='derived' else (COMB if row[9].value!='Established score' else None)
    if fill:
        for c in row: c.fill=fill
ws=wb.create_sheet('Comparison full text')
data=[]
for s,why in FT.items():
    p,doi=ids[s]
    data.append([int(s),p,doi,why,'Free in PMC ('+pmc[p]+')' if pmc.get(p) else 'NEED PDF',f'https://pubmed.ncbi.nlm.nih.gov/{p}/',''])
data.sort(key=lambda x:(x[4]!='NEED PDF',x[0]))
sheet(ws,['Serial (your file)','PMID','DOI','Comparison data missing from abstract','Full text','PubMed','Downloaded? (tick)'],data,[8,10,28,50,22,32,12])
wb.save(f)
print(len(rows),'result rows from',len({r['serial'] for r in rows}),'studies;',len(FT),'need full text, of which',sum(1 for p in pm if not pmc.get(p)),'not in PMC')
